#!/usr/bin/env python3
"""Dump a .docx file (CV or template) to plain text on stdout.

Usage:
    python3 read_docx.py path/to/file.docx

Prints paragraphs in order, then each table's cells row by row, so a CV's
work history is readable even when it's laid out in a table.
"""
import sys

from docx import Document


def dump_docx(path: str) -> str:
    doc = Document(path)
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for t_idx, table in enumerate(doc.tables):
        lines.append(f"\n--- table {t_idx} ---")
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit("usage: read_docx.py <file.docx>")
    print(dump_docx(sys.argv[1]))
