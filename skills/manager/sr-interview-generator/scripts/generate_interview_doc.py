#!/usr/bin/env python3
"""Fill the interview-notes docx template with generated questions.

Usage:
    python3 generate_interview_doc.py answers.json [output.docx]

answers.json shape:
{
  "name": "Jane Doe",
  "years_experience": "6",
  "categories": {
        "Exploration (HR Interview)": [
      {"context": "You've been at Cebu Pacific for 3 years —",
       "question": "what has changed in your role since you started?",
       "answer": "Concrete before/after responsibilities, scope growth, why now is the right time to move."},
      ... 2-3 entries
    ],
        "Technical Expertise (Java, Spring Boot, Backend Development)": [
      {"context": "You implemented testing with Karma/Jasmine at Cebu Pacific —",
       "question": "what was your strategy for testing AngularJS code that was hard to test?",
       "answer": "Concrete refactor for testability, mocking approach, coverage tradeoff."},
      ... 5 entries
    ],
        "Problem Solving, Debugging & Root Cause Analysis": [...],
        "Testing, Quality & Continuous Improvement": [...],
    "Agile Team Collaboration": [...],
        "Team Leading & Team Impact": [...]
  }
}

Preferred category keys match the human-facing labels in SKILL.md exactly.
Legacy short-form keys are still accepted for backwards compatibility. Each
category is normalized to the docx template placeholder it belongs to, then the
`[<Category> Questions]` placeholder paragraph is replaced with a
Question/Expected Answer table: the `context` clause is bold italic, the
`question` clause is highlighted yellow, and a bulleted notes row follows every
question row.
"""
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX

TEMPLATE = (
    Path(__file__).parent.parent
    / "template"
    / "Interview Notes - Senior Software Developer - [Candidate Name].docx"
)

DISPLAY_TO_TEMPLATE_CATEGORY = {
    "Exploration (HR Interview)": "Exploration",
    "Technical Expertise (Java, Spring Boot, Backend Development)": "Technical Expertise",
    "Problem Solving, Debugging & Root Cause Analysis": "Problem Solving",
    "Testing, Quality & Continuous Improvement": "Testing Quality",
    "Agile Team Collaboration": "Agile Team Collaboration",
    "Team Leading & Team Impact": "Team Leading",
}

LEGACY_CATEGORY_ALIASES = {
    "Exploration": "Exploration",
    "Technical Expertise": "Technical Expertise",
    "Problem Solving": "Problem Solving",
    "Testing Quality": "Testing Quality",
    "Agile Team Collaboration": "Agile Team Collaboration",
    "Team Leading": "Team Leading",
}

TEMPLATE_CATEGORY_ORDER = list(DISPLAY_TO_TEMPLATE_CATEGORY.values())
ACCEPTED_CATEGORY_ALIASES = {
    **DISPLAY_TO_TEMPLATE_CATEGORY,
    **LEGACY_CATEGORY_ALIASES,
}


def normalize_categories(raw_categories: dict[str, list]) -> list[tuple[str, list]]:
    normalized_categories: dict[str, list] = {}
    seen_as: dict[str, str] = {}

    for category_name, questions in raw_categories.items():
        template_category = ACCEPTED_CATEGORY_ALIASES.get(category_name)
        if template_category is None:
            accepted = ", ".join(sorted(ACCEPTED_CATEGORY_ALIASES))
            sys.exit(
                "unknown category key "
                f"{category_name!r}; expected one of: {accepted}"
            )
        if template_category in normalized_categories:
            sys.exit(
                "duplicate category alias for "
                f"{template_category!r}: {seen_as[template_category]!r} and {category_name!r}"
            )
        normalized_categories[template_category] = questions
        seen_as[template_category] = category_name

    missing = [
        display_name
        for display_name, template_category in DISPLAY_TO_TEMPLATE_CATEGORY.items()
        if template_category not in normalized_categories
    ]
    if missing:
        sys.exit(
            "missing required categories: "
            + ", ".join(missing)
        )

    return [
        (template_category, normalized_categories[template_category])
        for template_category in TEMPLATE_CATEGORY_ORDER
    ]


def replace_placeholder(doc: Document, placeholder: str, value: str) -> None:
    """Find `placeholder` in any paragraph (body or table cell) and swap it for `value`, keeping formatting."""
    def paragraphs():
        yield from doc.paragraphs
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from cell.paragraphs

    for p in paragraphs():
        if placeholder in p.text:
            full = p.text.replace(placeholder, value)
            for run in p.runs[1:]:
                run.text = ""
            if p.runs:
                p.runs[0].text = full
            else:
                p.add_run(full)


def add_question_row(table, context: str, question: str, answer: str) -> None:
    row = table.add_row()
    q_cell, a_cell = row.cells
    p = q_cell.paragraphs[0]
    r_context = p.add_run(context + " ")
    r_context.bold = True
    r_context.italic = True
    r_question = p.add_run(question)
    r_question.font.highlight_color = WD_COLOR_INDEX.YELLOW
    a_cell.paragraphs[0].add_run(answer)

    notes_row = table.add_row()
    notes_cell = notes_row.cells[0].merge(notes_row.cells[1])
    notes_cell.paragraphs[0].add_run("\u2022 ")


def build_category_table(doc: Document, placeholder_text: str, questions: list) -> None:
    target = next((p for p in doc.paragraphs if p.text.strip() == placeholder_text), None)
    if target is None:
        sys.exit(f"placeholder paragraph not found: {placeholder_text!r}")

    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].paragraphs[0].add_run("Question").bold = True
    hdr[1].paragraphs[0].add_run("Expected Answer").bold = True

    for q in questions:
        add_question_row(table, q["context"], q["question"], q["answer"])

    target._p.addnext(table._tbl)
    target._element.getparent().remove(target._element)


def generate(answers_path: str, output_path: str | None) -> str:
    answers = json.loads(Path(answers_path).read_text())
    doc = Document(TEMPLATE)
    categories = normalize_categories(answers["categories"])

    replace_placeholder(doc, "[Name]", answers["name"])
    replace_placeholder(doc, "[Career Experience]", str(answers["years_experience"]))

    for category, questions in categories:
        build_category_table(doc, f"[{category} Questions]", questions)

    if not output_path:
        output_path = str(
            TEMPLATE.parent.parent
            / f"Interview Notes - Senior Software Developer - {answers['name']}.docx"
        )
    doc.save(output_path)
    return output_path


if __name__ == "__main__":
    if len(sys.argv) not in (2, 3):
        sys.exit("usage: generate_interview_doc.py <answers.json> [output.docx]")
    out = generate(sys.argv[1], sys.argv[2] if len(sys.argv) == 3 else None)
    print(f"wrote {out}")
